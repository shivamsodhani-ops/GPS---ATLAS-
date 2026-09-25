"""Text extraction for every document type GPS actually deals with:
contracts/agreements (PDF/DOCX), BOQs/commercial terms (XLSX), drawings and
scanned approvals (images, OCR'd), MOMs/reports (DOCX/PDF/TXT), decks (PPTX).

Every extractor returns a list[PageText] so downstream chunking can keep
page-level provenance for citations ("Source: Vendor_Agreement.pdf, p.4").
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass


@dataclass
class PageText:
    page_number: int | None
    text: str


class ExtractionError(Exception):
    pass


def extract(filename: str, raw_bytes: bytes, mime_type: str) -> list[PageText]:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if ext == "pdf":
            return _extract_pdf(raw_bytes)
        if ext == "docx":
            return _extract_docx(raw_bytes)
        if ext in ("xlsx", "xlsm"):
            return _extract_xlsx(raw_bytes)
        if ext == "csv":
            return _extract_csv(raw_bytes)
        if ext in ("pptx",):
            return _extract_pptx(raw_bytes)
        if ext in ("txt", "md", "log"):
            return [PageText(1, raw_bytes.decode("utf-8", errors="ignore"))]
        if ext in ("png", "jpg", "jpeg", "tiff", "bmp"):
            return _extract_image_ocr(raw_bytes)
        # Legacy binary office formats (.doc, .xls, .ppt) -- best-effort via
        # LibreOffice headless conversion if it's installed; otherwise fail
        # loudly rather than silently indexing nothing.
        if ext in ("doc", "xls", "ppt", "rtf", "odt"):
            return _extract_via_libreoffice(filename, raw_bytes)
        raise ExtractionError(f"Unsupported file type: .{ext or 'unknown'}")
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 -- convert any parser crash into a clean status
        raise ExtractionError(f"Failed to parse .{ext} file: {exc}") from exc


# Render's free tier gives this whole app a small fraction of one shared
# CPU. OCR is by far the most expensive thing it does -- rendering a page to
# an image and running Tesseract over it can take several seconds of solid
# CPU time per page. A large scanned PDF run entirely inline (even in a
# background thread -- Python's GIL means CPU-bound work still competes with
# the main thread) can starve the process long enough that Render's health
# check (a 5-second timeout) fails and the instance gets killed and
# restarted, wiping the in-memory/ephemeral state of everyone's session.
# Capping both the render resolution and the page count bounds the worst
# case instead of letting one big scanned upload take the whole app down.
_OCR_RESOLUTION_DPI = 150
_OCR_MAX_PAGES = 25


def _extract_pdf(raw_bytes: bytes) -> list[PageText]:
    import pdfplumber

    pages: list[PageText] = []
    with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            if not text.strip():
                if i > _OCR_MAX_PAGES:
                    # Keep the page (and its citation numbering) but don't
                    # OCR it -- see the module note above. The document is
                    # still searchable on every page that had real text or
                    # was OCR'd; this only skips extra scanned pages past
                    # the cap.
                    text = ""
                else:
                    # scanned/image-only page -- fall back to OCR on a rendered image
                    try:
                        img = page.to_image(resolution=_OCR_RESOLUTION_DPI).original
                        text = _ocr_image(img)
                    except Exception:
                        text = ""
            pages.append(PageText(i, text))
    return pages


def _extract_docx(raw_bytes: bytes) -> list[PageText]:
    import docx

    document = docx.Document(io.BytesIO(raw_bytes))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    # DOCX has no reliable page boundaries without a rendering engine; treat
    # the whole document as one logical "page" for citation purposes.
    return [PageText(None, "\n".join(parts))]


def _extract_xlsx(raw_bytes: bytes) -> list[PageText]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True, read_only=True)
    pages = []
    for idx, sheet in enumerate(wb.worksheets, start=1):
        lines = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
        pages.append(PageText(idx, f"[Sheet: {sheet.title}]\n" + "\n".join(lines)))
    return pages


def _extract_csv(raw_bytes: bytes) -> list[PageText]:
    text = raw_bytes.decode("utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(text))
    lines = [" | ".join(row) for row in reader]
    return [PageText(1, "\n".join(lines))]


def _extract_pptx(raw_bytes: bytes) -> list[PageText]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(raw_bytes))
    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        lines.append(line)
        pages.append(PageText(i, "\n".join(lines)))
    return pages


def _ocr_image(pil_image) -> str:
    import pytesseract

    return pytesseract.image_to_string(pil_image)


def _extract_image_ocr(raw_bytes: bytes) -> list[PageText]:
    from PIL import Image

    img = Image.open(io.BytesIO(raw_bytes))
    text = _ocr_image(img)
    if not text.strip():
        raise ExtractionError("OCR produced no text -- image may be blank or unreadable")
    return [PageText(1, text)]


def _extract_via_libreoffice(filename: str, raw_bytes: bytes) -> list[PageText]:
    import shutil
    import subprocess
    import tempfile
    from pathlib import Path

    if shutil.which("soffice") is None:
        raise ExtractionError("Legacy office format needs LibreOffice, which is not installed on this server")

    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / filename
        src.write_bytes(raw_bytes)
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", td, str(src)],
            capture_output=True,
            timeout=90,
        )
        pdf_path = src.with_suffix(".pdf")
        if result.returncode != 0 or not pdf_path.exists():
            raise ExtractionError(f"LibreOffice conversion failed: {result.stderr.decode(errors='ignore')[:300]}")
        return _extract_pdf(pdf_path.read_bytes())
